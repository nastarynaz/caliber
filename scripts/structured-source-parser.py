#!/usr/bin/env python3
"""Deterministic XLSX/PPTX extraction with source locators.

This helper never modifies the source. JSON is written to stdout so the Node
ingestion coordinator can persist it transactionally in Supabase.
"""

from __future__ import annotations

import datetime as dt
import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET


def serializable(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def parse_xlsx(source: Path):
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required. Install scripts/requirements-ingestion.txt "
            "inside an isolated ingestion environment."
        ) from exc

    formulas = openpyxl.load_workbook(source, data_only=False, read_only=False)
    cached = openpyxl.load_workbook(source, data_only=True, read_only=False)
    pages = []
    try:
        for position, sheet in enumerate(formulas.worksheets, start=1):
            cached_sheet = cached[sheet.title]
            records = []
            readable_rows = []
            for row in sheet.iter_rows():
                cells = []
                readable = []
                for cell in row:
                    raw = cell.value
                    cached_value = cached_sheet[cell.coordinate].value
                    if raw is None and cached_value is None:
                        continue
                    formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                    display_value = serializable(cached_value if formula is not None and cached_value is not None else raw)
                    cells.append(
                        {
                            "cell": cell.coordinate,
                            "raw_value": serializable(raw),
                            "formula": formula,
                            "cached_value": serializable(cached_value),
                            "number_format": cell.number_format,
                            "data_type": cell.data_type,
                        }
                    )
                    readable.append(f"{cell.coordinate}={display_value}")
                if cells:
                    records.append({"row": row[0].row, "cells": cells})
                    readable_rows.append(" | ".join(readable))
            pages.append(
                {
                    "locator_type": "xlsx_range",
                    "locator_label": f"Sheet {position}: {sheet.title}",
                    "page_number": position,
                    "raw_text": "\n".join(readable_rows),
                    "regions": {
                        "sheet_name": sheet.title,
                        "sheet_index": position,
                        "max_row": sheet.max_row,
                        "max_column": sheet.max_column,
                        "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                        "rows": records,
                    },
                }
            )
    finally:
        formulas.close()
        cached.close()
    return {"kind": "xlsx", "pages": pages}


def slide_number(name: str):
    stem = Path(name).stem
    return int(stem.removeprefix("slide"))


def parse_pptx(source: Path):
    namespaces = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    }
    pages = []
    with zipfile.ZipFile(source) as archive:
        slide_names = sorted(
            (
                name
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ),
            key=slide_number,
        )
        media = sorted(name for name in archive.namelist() if name.startswith("ppt/media/"))
        for name in slide_names:
            number = slide_number(name)
            root = ET.fromstring(archive.read(name))
            paragraphs = []
            for paragraph in root.findall(".//a:p", namespaces):
                text = "".join(node.text or "" for node in paragraph.findall(".//a:t", namespaces)).strip()
                if text:
                    paragraphs.append(text)
            pages.append(
                {
                    "locator_type": "pptx_slide",
                    "locator_label": f"Slide {number}",
                    "page_number": number,
                    "raw_text": "\n".join(paragraphs),
                    "regions": {
                        "slide_entry": name,
                        "paragraph_count": len(paragraphs),
                        "package_media_entries": media,
                    },
                }
            )
    return {"kind": "pptx", "pages": pages}


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Usage: structured-source-parser.py <source.xlsx|source.pptx>")
    source = Path(sys.argv[1]).resolve(strict=True)
    suffix = source.suffix.lower()
    if suffix == ".xlsx":
        result = parse_xlsx(source)
    elif suffix == ".pptx":
        result = parse_pptx(source)
    else:
        raise SystemExit(f"Unsupported structured source: {suffix}")
    result.update({"source_filename": source.name, "schema_version": "structured-source-v1"})
    print(json.dumps(result, ensure_ascii=False, separators=(",", ":")))


if __name__ == "__main__":
    main()
